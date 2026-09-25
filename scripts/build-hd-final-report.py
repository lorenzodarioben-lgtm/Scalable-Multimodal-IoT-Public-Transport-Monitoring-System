"""Build the five-page HD report and separate evidence appendix from reviewed data."""

from pathlib import Path
import json
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION_START
from docx.shared import Cm, Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "hd-final-report"
OUT.mkdir(parents=True, exist_ok=True)
AGG = json.loads((ROOT / "artifacts/hd-analysis/aggregate.json").read_text())
RUNS = json.loads((ROOT / "artifacts/hd-analysis/run-metrics.json").read_text())
FIG = ROOT / "docs/hd-final-figures"


def style(doc, landscape=False):
    sec = doc.sections[0]
    sec.page_width, sec.page_height = ((Cm(29.7), Cm(21)) if landscape else (Cm(21), Cm(29.7)))
    sec.top_margin = Cm(1.65)
    sec.bottom_margin = Cm(1.45)
    sec.left_margin = Cm(1.8)
    sec.right_margin = Cm(1.8)
    sec.header_distance = Cm(0.7)
    sec.footer_distance = Cm(0.7)
    normal = doc.styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(9.5)
    normal.font.color.rgb = RGBColor(20, 24, 28)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.10
    for name, size, before, after in [("Title", 16, 0, 7), ("Heading 1", 11.5, 8, 4),
                                      ("Heading 2", 10, 6, 3)]:
        s = doc.styles[name]
        s.font.name = "Aptos"
        s.font.size = Pt(size)
        s.font.bold = True
        s.font.color.rgb = RGBColor(20, 24, 28)
        s.paragraph_format.space_before = Pt(before)
        s.paragraph_format.space_after = Pt(after)
        s.paragraph_format.keep_with_next = True
        if name == "Title":
            borders = s.element.get_or_add_pPr().find(qn("w:pBdr"))
            if borders is not None:
                s.element.get_or_add_pPr().remove(borders)
    footer = sec.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.style = normal
    footer.add_run("Lorenzo Dario Ben | s224658462 | SIT314 6.4HD | ")
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)


def p(doc, text, bold_prefix=None):
    para = doc.add_paragraph()
    if bold_prefix and text.startswith(bold_prefix):
        para.add_run(bold_prefix).bold = True
        para.add_run(text[len(bold_prefix):])
    else:
        para.add_run(text)
    return para


def heading(doc, text):
    doc.add_paragraph(text, "Heading 1")


def caption(doc, text):
    para = doc.add_paragraph(text)
    para.style = doc.styles["Normal"]
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    para.paragraph_format.space_after = Pt(5)
    for run in para.runs:
        run.italic = True
        run.font.size = Pt(8)
    return para


def figure(doc, name, width=16.4):
    src = FIG / f"{name}.png"
    if not src.exists():
        raise FileNotFoundError(f"Render the formal SVG to PNG first: {src}")
    para = doc.add_paragraph()
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    para.paragraph_format.space_after = Pt(0)
    para.add_run().add_picture(str(src), width=Cm(width))


def table(doc, rows, widths=None, font=8.1):
    t = doc.add_table(rows=0, cols=len(rows[0]))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.style = "Table Grid"
    t.autofit = False
    for ridx, row in enumerate(rows):
        cells = t.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = str(val)
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if widths:
                cells[i].width = Cm(widths[i])
            for para in cells[i].paragraphs:
                para.paragraph_format.space_after = Pt(1)
                para.paragraph_format.space_before = Pt(1)
                para.paragraph_format.line_spacing = 1
                for run in para.runs:
                    run.font.size = Pt(font)
                    if ridx == 0:
                        run.bold = True
        if ridx == 0:
            for cell in cells:
                tcpr = cell._tc.get_or_add_tcPr()
                shd = OxmlElement("w:shd")
                shd.set(qn("w:fill"), "E9ECEF")
                tcpr.append(shd)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return t


def page(doc):
    doc.add_page_break()


def fmt(value, digits=1):
    if value is None:
        return "n/a"
    return f"{value:.{digits}f}" if isinstance(value, (int, float)) else str(value)


def group(workload, arm, metric):
    return AGG["groups"][workload][arm]["metrics"][metric]


def result_table(doc, workload):
    specs = [
        ("First request from onset, s", "scaleRequestLatencySeconds", 1),
        ("Request to first ready, s", "requestToFirstReadySeconds", 1),
        ("Peak visible jobs", "peakVisibleBacklog", 0),
        ("Peak BPT", "peakBacklogPerTask", 1),
        ("Oldest message, s", "peakOldestMessageAgeSeconds", 1),
        ("Throughput, jobs/s", "completionThroughputJobsPerSecond", 2),
        ("Drain, s", "drainSeconds", 1),
        ("Processing p50, ms", "processingP50Ms", 1),
        ("Processing p95, ms", "processingP95Ms", 1),
        ("Peak running tasks", "peakRunningTasks", 0),
        ("Task-seconds", "taskSeconds", 0),
    ]
    rows = [["Outcome", "Reactive r1/r2/r3", "Mean ± SD", "Hybrid r1/r2/r3", "Mean ± SD"]]
    for label, key, digits in specs:
        r, h = group(workload, "reactive", key), group(workload, "hybrid", key)
        rows.append([label, "/".join(fmt(v, digits) for v in r["raw"]),
                     f"{fmt(r['mean'], digits)} ± {fmt(r['sampleSd'], digits)}",
                     "/".join(fmt(v, digits) for v in h["raw"]),
                     f"{fmt(h['mean'], digits)} ± {fmt(h['sampleSd'], digits)}"])
    table(doc, rows, [3.7, 3.4, 2.4, 3.4, 2.4], 7.5)


main = Document()
style(main)
main.add_paragraph("Hybrid predictive and reactive autoscaling for a bursty IoT pipeline", "Title")
p(main, "SIT314 6.4HD final research report | Lorenzo Dario Ben | s224658462 | 26 September 2026")
heading(main, "Abstract")
p(main, "A separately deployed AWS IoT analysis pipeline was evaluated under predictable-ramp and sudden-burst traffic. We compared a fast reactive backlog alarm with an otherwise identical hybrid controller that forecasts analysis-job arrivals and can request additional workers. All 12 planned runs passed manual validity review. Across three ramp repeats, hybrid scale requests led the declared high-load onset by 61.3 s on average and cut mean peak BacklogPerTask (BPT) from 796.3 to 43.0, but used 66.5% more running task-seconds. In burst repeats, hybrid requests came 23.8 s after onset rather than 100.9 s for reactive control; mean peak BPT fell from 766.3 to 222.8, with 17.5% more task-seconds. All 208,800 submitted analysis jobs completed without accounting or reliability faults. The result supports an earlier, workload-specific scale-out benefit, not a claim that a short-horizon model predicts an unannounced burst or reduces cost.")
heading(main, "Research question and context")
p(main, "Can a lightweight analysis-arrival predictor, added to unchanged reactive safeguards, reduce scale-out delay and queue pressure during a learnable traffic ramp while remaining safe under an abrupt burst? The prior Distinction system used SQS-backed ECS Fargate workers, target tracking at 75 queued jobs per task, and a fast 60 s alarm that added four tasks when BPT exceeded 75. In its final improvement retest, the alarm still took 62.512 s from first above-target minute to scale request and queue depth reached 1,042. This is background motivation, not a matched HD control: the matched comparison is reactive versus hybrid within the isolated HD deployment.")
heading(main, "Research literature")
p(main, "Wang, Chandra and Weissman's Jingle uses IoT-derived context in a hybrid predictive-reactive edge autoscaler [1]. Masdari and Khoshnevis classify forecasting approaches for proactive cloud resource management [2]. Kumar, Goomer and Singh evaluate an LSTM cloud-workload forecaster [3]. These studies motivate a domain-proximate arrival signal and a predictive comparison, but do not establish that their models transfer to this pipeline. An auditable linear trend was chosen here instead of training a neural model on a small synthetic workload. We did not reproduce or benchmark Jingle or LSTM.")
heading(main, "Contribution")
p(main, "The contribution is a frozen, explainable scale-out-only hybrid path, plus a matched AWS evaluation that reports both pressure reduction and capacity-time cost. Distinction production resources and evidence were not changed.")

page(main)
heading(main, "System and frozen controller")
p(main, "The HD stack has isolated SQS analysis, arrival and notification queues, ECS workers, DynamoDB predictor state and SIT314/HDTransport metrics. After successful analysis-job fan-out, an idempotent run-tagged observation reports published jobs. A Lambda groups arrivals into 10 s bins and fits ordinary least squares to eight chronological rates. Its 80 s forecast is max(0, a + b(tnow + 80)). With measured single-task capacity c = 42.467 jobs/s, queue B, running/desired/reactive floor n and horizon H = 80 s, it estimates Bfuture = B + max(0, forecast - nc)H and recommends clamp(1,5,max(n,ceil(forecast/c),ceil(Bfuture/75))). A slope of at least 0.02 jobs/s², two consecutive positive recommendations and a 60 s suppression of same-or-lower duplicate requests prevent noisy action. The controller can request scale-out only. Original BPT target tracking, target 75 and fast +4 alarm remain active in both arms; no alarms were forced. Parameters were fixed before formal AWS runs after local sensitivity analysis, which is not counted as cloud evidence.")
heading(main, "Matched AWS method and validity")
p(main, "Each workload has a 30 s warm-up and 600 s measurement window. The predictable ramp rises through 10, 16.667, 25, 33.333 and 50 jobs/s for 330 incidents and 16,500 analysis jobs. The sudden burst stays at 10 jobs/s, jumps to 50 jobs/s from 210 to 510 s, and returns to 10 jobs/s for 366 incidents and 18,300 jobs. Incident payloads, paired logical digests, 50 ms worker processing delay, one-to-five task bound, scaling policies and timing guard are identical across arms. Three reactive and three hybrid repeats were run for each class, each with fresh execution identity and a clean one-task start. Only the HD controller mode changed between arms.")
p(main, "Manual reviews checked the unchanged schedule-lag gate, published/completed counts, duplicate and DLQ accounting, genuine historical CloudWatch BPT, scale activities, task RUNNING and WORKER_READY logs and paired workload digests. SQS visible depth and age are sampled observations, not continuous maxima. Task-seconds integrate sampled running tasks over the common 600 s measurement interval and indicate relative capacity use, not billed AWS cost. Medians, sample SDs and every raw repeat are in the evidence appendix and CSV. Means are descriptive; n = 3 does not support a significance claim.")
heading(main, "Excluded infrastructure attempt")
p(main, "The first sudden-burst hybrid r2 attempt stopped at 60/366 incidents and 3,000/18,300 jobs after a DNS/SQS endpoint resolution failure. Its raw artifact and INVALID review are preserved but excluded from all formal estimates. A new, independently reviewed replacement r2 completed the same frozen logical workload and is the only valid hybrid r2 in the 12-run aggregate. No valid run was reinjected.")

page(main)
heading(main, "Predictable ramp results")
p(main, "All six ramp runs completed 16,500/16,500 jobs; offered rates were close (reactive mean 26.177, hybrid 26.055 jobs/s). The first hybrid request preceded the declared 510 s high-load onset in every repeat (45.9, 76.3 and 61.6 s early); reactive requests followed it by 88.7, 86.6 and 76.3 s. The hybrid's first new worker became ready 26.6, 41.5 and 21.0 s after request, while the reactive path took 32.4, 24.1 and 22.8 s. The gain therefore comes from earlier requests, not faster Fargate startup.")
result_table(main, "PREDICTABLE_RAMP")
figure(main, "figure-1-ramp-peak-bpt", 14.4)
caption(main, "Figure 1. Genuine historical CloudWatch peak BPT for every matched ramp repeat. The target is 75.")
p(main, "Mean peak visible queue fell 94.3%, genuine BPT 94.6%, and oldest-message age 40.0%. However, mean task-seconds rose from 633.952 to 1055.815 (+66.5%). Throughput changed by -0.3%, mean drain by +0.45 s, and mean p95 processing latency from 52.3 to 53.0 ms. Prediction MAE was 8.730 jobs/s with signed bias -1.496; no false proactive requests were recorded. Figure 4 and its CSV show the representative r2 rate, request, readiness and overload chronology.")

page(main)
heading(main, "Sudden burst results")
p(main, "All six burst runs completed 18,300/18,300 jobs at similar offered rates (reactive mean 29.170, hybrid 29.219 jobs/s). An unannounced jump at 210 s cannot be forecast before it appears in the history. Indeed, none of the three hybrid scale requests preceded onset. They occurred 23.3, 24.6 and 23.3 s afterward; reactive requests followed at 110.6, 99.3 and 92.7 s. A predictor action happened earlier after the jump, but should not be described as pre-burst foresight. The unchanged reactive alarm remained a safety path.")
result_table(main, "SUDDEN_BURST")
figure(main, "figure-5-burst-peak-bpt", 14.4)
caption(main, "Figure 5. Genuine historical CloudWatch peak BPT for every matched burst repeat.")
p(main, "Mean peak visible queue fell 50.9%, genuine BPT 70.9%, and oldest-message age 34.0%. Mean task-seconds rose from 1759.971 to 2067.700 (+17.5%). Throughput changed by -0.7%, mean drain increased by 4.19 s, and mean p95 latency by 1.0 ms. Burst forecast MAE was 22.470 jobs/s, materially higher than for the ramp, with signed bias +6.270. No false proactive request was logged under the frozen review definition.")

page(main)
heading(main, "Interpretation and limitations")
p(main, "The ramp result answers the question positively within this workload: the hybrid requested workers before overload in all three repeats, and queue pressure was lower in every matched pair. For the sudden burst, it reacted earlier after onset and reduced pressure in every pair, but did not predict the discontinuity. The first-ready delay after a request remained roughly 26–30 s across arms, so advance notice or a faster post-onset trigger, rather than a change in worker startup, explains the observed difference. Reliability was unchanged: all 208,800 formal jobs completed, with zero failures, duplicate results, DLQ messages or unaccounted jobs. Each arm reached up to the five-task ceiling in burst repeats; peak count alone does not show how long capacity was used.")
p(main, "The cost of lower pressure is real in the sampled capacity proxy: +66.5% task-seconds in the ramp and +17.5% in the burst. Mean completion throughput did not improve, and hybrid drain and p95 latency were slightly worse. This study therefore supports pressure reduction at a capacity-time premium, not a universal performance or financial win. It does not calculate AWS charges. The small three-repeat sample, one region, synthetic deterministic shapes, chosen 80 s horizon, 1–5 task range, sampled queue telemetry and CloudWatch publication granularity limit generalisation. The local tuning grid may fit the designed ramp. Worker startup varied across repeats, and the DNS failure illustrates that infrastructure faults can invalidate an attempt independently of the controller. No statistical significance or production SLO guarantee is claimed.")
heading(main, "Conclusion")
p(main, "A lightweight arrival-rate trend provided useful lead time for a predictable ramp and earlier post-onset scale-out for an abrupt burst. In this isolated AWS experiment it reduced queue pressure consistently without job-accounting faults, but consumed more running task time and did not improve throughput or drain. Keep the reactive safeguards, and treat forecast quality and capacity-time as explicit design constraints before production adoption. The complete review decisions, raw run IDs, summary statistics, CSV sources and six reproducible figures are supplied in the separate evidence appendix.")
heading(main, "References")
refs = [
    "[1] Y. Wang, A. Chandra and J. Weissman, ‘Jingle: IoT-Informed Autoscaling for Efficient Resource Management in Edge Computing,’ CCGrid, pp. 395–407, 2024. doi:10.1109/CCGrid59990.2024.00052.",
    "[2] M. Masdari and A. Khoshnevis, ‘A survey and classification of the workload forecasting methods in cloud computing,’ Cluster Computing, 23, pp. 2399–2424, 2020. doi:10.1007/s10586-019-03010-3.",
    "[3] J. Kumar, R. Goomer and A. K. Singh, ‘Long Short Term Memory Recurrent Neural Network (LSTM-RNN) Based Workload Forecasting Model For Cloud Datacenters,’ Procedia Computer Science, 125, pp. 676–682, 2018. doi:10.1016/j.procs.2017.12.087.",
]
for ref in refs:
    q = p(main, ref)
    q.paragraph_format.space_after = Pt(2)
    for run in q.runs:
        run.font.size = Pt(8)
main.save(OUT / "SIT314_6_4HD_Final_Report.docx")


appendix = Document()
style(appendix, landscape=True)
appendix.add_paragraph("SIT314 6 4HD formal evidence appendix", "Title")
p(appendix, "Lorenzo Dario Ben | s224658462 | 12 VALID AWS runs, one preserved INVALID infrastructure attempt")
heading(appendix, "Evidence rules and provenance")
p(appendix, "The aggregate includes exactly three VALID reactive and three VALID hybrid runs for each of two workload classes. The partial DNS/SQS failure is retained in artifacts/hd-aws-runs but excluded. Each valid row has raw samples, summary, historical CloudWatch data, scaling and worker-readiness evidence and a manual review. Peak BPT comes from the SIT314/HDTransport CloudWatch metric, not a calculation from queue depth. Complete machine-readable run records and aggregate mean, median, sample SD, absolute and percentage differences are in docs/hd-final-data/all-reviewed-runs.csv and all-aggregate-statistics.csv. Figure CSVs are in docs/hd-final-data and the six SVGs are in docs/hd-final-figures. The local raw files are indexed with SHA-256 in docs/hd-final-data/raw-evidence-sha256.csv.")
heading(appendix, "Twelve reviewed run identifiers")
rows = [["Class", "Arm", "r", "Run identifier", "Jobs", "Digest prefix"]]
for run in RUNS:
    rows.append(["Ramp" if run["workloadClass"] == "PREDICTABLE_RAMP" else "Burst",
                 run["arm"], run["repeatNumber"], run["runId"],
                 run["completedJobs"], run["logicalDigest"][:12]])
table(appendix, rows, [1.3, 1.4, 0.6, 17.0, 1.2, 2.5], 6.7)
page(appendix)
heading(appendix, "Per-run primary outcomes")
rows = [["Class", "Arm", "r", "Request s", "Ready s", "Visible", "BPT", "Age s", "Task-s", "Rate", "Drain s", "p95 ms", "Faults"]]
for run in RUNS:
    rows.append(["Ramp" if run["workloadClass"] == "PREDICTABLE_RAMP" else "Burst",
                 run["arm"], run["repeatNumber"], fmt(run["scaleRequestLatencySeconds"], 1),
                 fmt(run["requestToFirstReadySeconds"], 1), fmt(run["peakVisibleBacklog"], 0),
                 fmt(run["peakBacklogPerTask"], 1), fmt(run["peakOldestMessageAgeSeconds"], 0),
                 fmt(run["taskSeconds"], 1), fmt(run["completionThroughputJobsPerSecond"], 2),
                 fmt(run["drainSeconds"], 2), fmt(run["processingP95Ms"], 0),
                 run["errors"] + run["duplicates"] + run["dlq"] + run["unaccountedJobs"]])
table(appendix, rows, [1.4, 1.5, 0.7, 2, 1.8, 1.7, 1.7, 1.5, 1.7, 1.6, 1.6, 1.4, 1.2], 7.3)
p(appendix, "Request s is relative to declared high-load onset; negative ramp values are proactive lead. Ready s is request to first new WORKER_READY. Rate is completion throughput in jobs/s. All 12 valid runs submitted and completed their full workload. Faults combines errors, duplicate results, DLQ and unaccounted jobs; all are zero. The complete CSV retains each component separately, plus schedule lag, peak tasks, prediction error and paired digest.")
heading(appendix, "Aggregate means, medians and sample standard deviations")
rows = [["Class", "Metric", "Reactive mean", "Reactive median", "Reactive SD", "Hybrid mean", "Hybrid median", "Hybrid SD", "Δ mean", "Δ %"]]
keys = [("offeredJobsPerSecond", "Offered jobs/s"),
        ("scaleRequestLatencySeconds", "Request from onset s"),
        ("requestToFirstReadySeconds", "Request to ready s"),
        ("peakVisibleBacklog", "Visible jobs"),
        ("peakBacklogPerTask", "BPT"), ("peakOldestMessageAgeSeconds", "Age s"),
        ("completionThroughputJobsPerSecond", "Throughput"), ("drainSeconds", "Drain s"),
        ("processingP50Ms", "p50 ms"), ("processingP95Ms", "p95 ms"),
        ("peakRunningTasks", "Peak tasks"), ("taskSeconds", "Task-s")]
for wl, label in [("PREDICTABLE_RAMP", "Ramp"), ("SUDDEN_BURST", "Burst")]:
    for key, metric_label in keys:
        r, h = group(wl, "reactive", key), group(wl, "hybrid", key)
        delta = h["mean"] - r["mean"]
        pct = delta / r["mean"] * 100 if r["mean"] else None
        rows.append([label, metric_label, fmt(r["mean"], 2), fmt(r["median"], 2),
                     fmt(r["sampleSd"], 2), fmt(h["mean"], 2), fmt(h["median"], 2),
                     fmt(h["sampleSd"], 2), fmt(delta, 2), fmt(pct, 1) + "%" if pct is not None else "n/a"])
table(appendix, rows, [1.3, 3.4, 2.3, 2.3, 2, 2.3, 2.3, 2, 2, 1.6], 7.0)
p(appendix, "The ramp request-timing percentage crosses zero and is only an arithmetic signed-time change. The interpretable result is a mean request 145.142 s earlier, switching from 83.883 s after onset to 61.259 s before onset.")

heading(appendix, "Figure and data index")
for num, title in enumerate([
    "Predictable ramp peak BPT by repeat", "Predictable ramp mean queue pressure",
    "Predictable ramp task-seconds by repeat", "Predictable ramp r2 arrival and scaling timeline",
    "Sudden burst peak BPT by repeat", "Sudden burst task-seconds by repeat",
], 1):
    p(appendix, f"Figure {num}: {title}. SVG: docs/hd-final-figures/figure-{num}-*.svg. Source CSV: docs/hd-final-data/figure-{num}-*.csv.")
heading(appendix, "Interpretation checks")
p(appendix, "The ramp's hybrid requests are all before 510 s high-load onset, whereas the sudden-burst hybrid requests are all after 210 s onset. The 80 s forecast does not anticipate an unknown step change. Request-to-first-ready time is similar across arms, so earlier request timing is the observed mediator. Every class/arm has three reviewed repeats; all values, including higher task-seconds, slightly lower throughput, longer hybrid mean drain and marginally higher p95 latency, remain visible. The data do not demonstrate lower AWS cost or statistical significance.")
heading(appendix, "Excluded infrastructure attempt")
p(appendix, "2026-09-23T23-06-57-400Z-hd-sudden-burst-hybrid-r2-22296a14: INVALID, 60/366 incidents and 3,000/18,300 jobs after DNS/SQS endpoint resolution failure. Replacement run a07532a9 is VALID and shares reactive r2's frozen logical digest. The invalid raw evidence and review are preserved but not aggregated.")
heading(appendix, "Forecast error and reliability")
rows = [["Class", "Hybrid MAE r1/r2/r3", "Mean ± sample SD", "Hybrid bias r1/r2/r3", "Mean ± sample SD"]]
for wl, label in [("PREDICTABLE_RAMP", "Ramp"), ("SUDDEN_BURST", "Burst")]:
    mae, bias = group(wl, "hybrid", "predictionMaeJobsPerSecond"), group(wl, "hybrid", "predictionBiasJobsPerSecond")
    rows.append([label, "/".join(fmt(x, 2) for x in mae["raw"]),
                 f"{fmt(mae['mean'], 2)} ± {fmt(mae['sampleSd'], 2)}",
                 "/".join(fmt(x, 2) for x in bias["raw"]),
                 f"{fmt(bias['mean'], 2)} ± {fmt(bias['sampleSd'], 2)}"])
table(appendix, rows, [2, 5, 4, 5, 4], 7.5)
p(appendix, "All 12 valid runs recorded zero errors, duplicate results, duplicate jobs skipped, DLQ messages and unaccounted jobs. All 12 passed the frozen schedule-lag validity guard. The underlying per-run CSV retains every schedule-lag and reliability field.")
heading(appendix, "Final operational state at 26 September 2026 09:21:36 UTC")
p(appendix, "HD analysis, arrival and notification queues and their DLQs were each 0 visible / 0 in flight. HD route ECS was desired/running/pending 1/1/0 with exact current WORKER_READY at 08:16:11.488 UTC; notification ECS was 1/1/0. Recent genuine HD BPT was 0; fast alarm OK; predictor hybrid mode had no pending request; target tracking stayed at 75 with 1–5 capacity and fast +4 intact. Distinction route ECS remained 1/1/0, its analysis queue/DLQ 0/0, BPT target 75 and fast +4 intact, recent BPT 0. AWS infrastructure was not cleaned up and must remain for review.")
appendix.save(OUT / "SIT314_6_4HD_Evidence_Appendix.docx")
print("Created report and appendix DOCX from 12 reviewed runs")
